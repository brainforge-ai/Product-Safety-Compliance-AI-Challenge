"""Extract text from TXT, JSON, and DOCX files."""

import json


def extract_from_text(file_path: str) -> str:
    """Read a plain-text file and return its contents."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def extract_from_json(file_path: str) -> str:
    """Read JSON and convert to a readable text representation."""
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    if isinstance(data, dict):
        return _dict_to_text(data)
    if isinstance(data, list):
        return "\n\n".join(_dict_to_text(item) if isinstance(item, dict) else str(item) for item in data)
    return str(data)


def extract_from_docx(file_path: str) -> str:
    """Read a DOCX file, extracting both paragraph text and table contents."""
    from docx import Document
    doc = Document(file_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _dict_to_text(d: dict) -> str:
    """Flatten a dict into key: value lines for the ingredient parser."""
    lines = []
    for key, value in d.items():
        if isinstance(value, list):
            lines.append(f"{key}:")
            for item in value:
                lines.append(f"- {item}")
        else:
            lines.append(f"{key}: {value}")
    return "\n".join(lines)
